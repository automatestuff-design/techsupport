import base64
import io
import json
import os
import re
import uuid
from pathlib import Path

import httpx
from anthropic import AsyncAnthropic


def _get_client() -> AsyncAnthropic:
    key = os.environ.get("ANTHROPIC_API_KEY")
    if not key:
        raise RuntimeError("ANTHROPIC_API_KEY is not set.")
    return AsyncAnthropic(api_key=key)


def extract_text_from_pdf(data: bytes) -> str:
    import pdfplumber
    parts = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text and text.strip():
                parts.append(text.strip())
    return "\n\n".join(parts)


def extract_text_from_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                parts.append(row_text)
    return "\n\n".join(parts)


async def extract_from_image(data: bytes, media_type: str, filename: str) -> dict:
    """Use Claude vision to OCR and describe an image."""
    b64 = base64.standard_b64encode(data).decode()
    response = await _get_client().messages.create(
        model="claude-opus-4-7",
        max_tokens=2048,
        messages=[{
            "role": "user",
            "content": [
                {
                    "type": "image",
                    "source": {"type": "base64", "media_type": media_type, "data": b64},
                },
                {
                    "type": "text",
                    "text": (
                        f'Analyze this image (filename: "{filename}").\n'
                        "Extract ALL visible text. If it shows a technical support scenario, "
                        "error message, screenshot, or document, describe it in full detail.\n\n"
                        "Reply as JSON only:\n"
                        '{"title": "concise descriptive title (max 80 chars)", '
                        '"content": "full extracted text and description", '
                        '"category": "one of: Connectivity/Email/Printing/Performance/'
                        'Authentication/Hardware/Software/Security/Other"}'
                    ),
                },
            ],
        }],
    )
    raw = response.content[0].text.strip()
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", raw, re.DOTALL)
        if match:
            try:
                return json.loads(match.group())
            except json.JSONDecodeError:
                pass
    stem = filename.rsplit(".", 1)[0].replace("_", " ").replace("-", " ").title()
    return {"title": stem, "content": raw, "category": None}


async def upload_to_supabase_storage(data: bytes, filename: str, content_type: str) -> str | None:
    """Upload a file to Supabase Storage and return its public URL, or None on failure."""
    supabase_url = os.environ.get("SUPABASE_URL", "")
    service_role_key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "")
    if not supabase_url or not service_role_key:
        return None

    bucket = "transcripts"
    stem = Path(filename).stem[:50].replace(" ", "_")
    ext = Path(filename).suffix.lower()
    storage_path = f"uploads/{stem}_{uuid.uuid4().hex[:8]}{ext}"

    async with httpx.AsyncClient() as client:
        resp = await client.post(
            f"{supabase_url}/storage/v1/object/{bucket}/{storage_path}",
            content=data,
            headers={
                "Authorization": f"Bearer {service_role_key}",
                "apikey": service_role_key,
                "Content-Type": content_type,
            },
            timeout=60.0,
        )

    if resp.status_code not in (200, 201):
        return None

    return f"{supabase_url}/storage/v1/object/public/{bucket}/{storage_path}"


async def generate_metadata(text: str, filename: str) -> dict:
    """Ask Claude to suggest a title and category for extracted document text."""
    stem = filename.rsplit(".", 1)[0].replace("_", " ").replace("-", " ").title()
    response = await _get_client().messages.create(
        model="claude-opus-4-7",
        max_tokens=128,
        messages=[{
            "role": "user",
            "content": (
                f'Filename: "{stem}"\n\nDocument excerpt:\n{text[:3000]}\n\n'
                "Provide a concise title (max 80 chars) and a category.\n"
                "Categories: Connectivity, Email, Printing, Performance, Authentication, "
                "Hardware, Software, Security, Other\n\n"
                'Reply as JSON only: {"title": "...", "category": "..."}'
            ),
        }],
    )
    raw = response.content[0].text.strip()
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", raw, re.DOTALL)
        if match:
            try:
                return json.loads(match.group())
            except json.JSONDecodeError:
                pass
    return {"title": stem, "category": None}
